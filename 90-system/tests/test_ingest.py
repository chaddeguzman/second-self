from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from second_self.ingest import ingest
from second_self.paths import SecondSelfPaths


def test_txt_ingest_is_provenance_safe(second_self: SecondSelfPaths, tmp_path: Path) -> None:
    source = tmp_path / "journal.txt"
    source.write_text("A journal entry.", encoding="utf-8")
    first = ingest(second_self, source)
    second = ingest(second_self, source)
    assert first["duplicate"] is False
    assert second["duplicate"] is True
    assert Path(str(first["original"])).read_text(encoding="utf-8") == "A journal entry."
    assert "A journal entry." in Path(str(first["extracted"])).read_text(encoding="utf-8")


def test_docx_and_xlsx_extraction(second_self: SecondSelfPaths, tmp_path: Path) -> None:
    docx_path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("Purpose note")
    document.save(docx_path)
    docx_result = ingest(second_self, docx_path)
    assert "Purpose note" in Path(str(docx_result["extracted"])).read_text(encoding="utf-8")

    xlsx_path = tmp_path / "lessons.xlsx"
    workbook = Workbook()
    workbook.active.title = "Lessons"
    workbook.active.append(["Lesson", "Source"])
    workbook.active.append(["Keep provenance", "Review"])
    workbook.save(xlsx_path)
    xlsx_result = ingest(second_self, xlsx_path)
    extracted = Path(str(xlsx_result["extracted"])).read_text(encoding="utf-8")
    assert "Sheet: Lessons" in extracted
    assert "Keep provenance | Review" in extracted


def test_image_only_pdf_is_flagged(second_self: SecondSelfPaths, tmp_path: Path) -> None:
    pdf_path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with pdf_path.open("wb") as stream:
        writer.write(stream)
    result = ingest(second_self, pdf_path)
    assert any("scanned" in warning for warning in result["warnings"])

